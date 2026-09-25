"""Resume file parsing (PDF / DOCX) and text cleaning."""
from __future__ import annotations

import io
import logging
import re
import unicodedata
from dataclasses import dataclass

import docx  # python-docx
import fitz  # PyMuPDF

logger = logging.getLogger(__name__)

SUPPORTED_TYPES = {"pdf", "docx"}


class ResumeParseError(ValueError):
    """Raised when a resume cannot be parsed into usable text."""


@dataclass
class ParsedDocument:
    text: str
    file_type: str
    pages: int


def _extract_pdf(data: bytes) -> tuple[str, int]:
    try:
        with fitz.open(stream=data, filetype="pdf") as doc:
            if doc.needs_pass:
                raise ResumeParseError("PDF is password protected.")
            pages = [page.get_text("text", sort=True) for page in doc]
            return "\n".join(pages), len(pages)
    except ResumeParseError:
        raise
    except Exception as exc:  # PyMuPDF raises a variety of runtime errors on corrupt input
        raise ResumeParseError(f"Could not read PDF: {exc}") from exc


def _extract_docx(data: bytes) -> tuple[str, int]:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ResumeParseError(f"Could not read DOCX: {exc}") from exc
    parts: list[str] = [p.text for p in document.paragraphs]
    # Many resume templates put content in tables (two-column layouts).
    for table in document.tables:
        for row in table.rows:
            seen: set[int] = set()
            for cell in row.cells:
                if id(cell._tc) in seen:  # merged cells repeat
                    continue
                seen.add(id(cell._tc))
                parts.append(cell.text)
    return "\n".join(parts), 1


_BULLETS = re.compile(r"^[ \t]*[•●▪◦■□►▸‣⁃∙·\*o]\s+", re.MULTILINE)
_DASH_BULLET = re.compile(r"^[ \t]*[–—-]\s+", re.MULTILINE)
_HYPHEN_BREAK = re.compile(r"(\w)-\n(\w)")
_MULTISPACE = re.compile(r"[ \t ]+")
_MULTINEWLINE = re.compile(r"\n{3,}")


def clean_text(text: str) -> str:
    """Normalize unicode, bullets and whitespace while preserving line structure (needed for sections)."""
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = "".join(ch for ch in text if ch in "\n\t" or unicodedata.category(ch)[0] != "C")
    text = _HYPHEN_BREAK.sub(r"\1\2", text)
    text = _BULLETS.sub("- ", text)
    text = _DASH_BULLET.sub("- ", text)
    text = text.replace("–", "-").replace("—", "-")
    text = _MULTISPACE.sub(" ", text)
    lines = [ln.strip() for ln in text.split("\n")]
    text = "\n".join(lines)
    return _MULTINEWLINE.sub("\n\n", text).strip()


def parse_resume(filename: str, data: bytes) -> ParsedDocument:
    """Extract and clean text from an uploaded resume."""
    if not data:
        raise ResumeParseError("Uploaded file is empty.")
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in SUPPORTED_TYPES:
        raise ResumeParseError(f"Unsupported file type '.{ext}'. Upload a PDF or DOCX resume.")
    raw, pages = _extract_pdf(data) if ext == "pdf" else _extract_docx(data)
    text = clean_text(raw)
    if len(text.split()) < 20:
        raise ResumeParseError(
            "Very little text could be extracted. If this is a scanned/image PDF, export it as a text-based PDF or DOCX."
        )
    logger.info("Parsed %s (%s, %d pages, %d words)", filename, ext, pages, len(text.split()))
    return ParsedDocument(text=text, file_type=ext, pages=pages)
